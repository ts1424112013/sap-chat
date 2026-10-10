from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]


def mark(run):
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def add_p(doc, text, style=None, bold=False):
    if style and style not in [s.name for s in doc.styles]:
        style = None
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    r.bold = bold
    mark(r)
    return p


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if "Table Grid" in [s.name for s in doc.styles]:
        try:
            t.style = "Table Grid"
        except KeyError:
            pass
    for cell, value in zip(t.rows[0].cells, headers):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        r = cell.paragraphs[0].add_run(value)
        r.bold = True
        mark(r)
    for row in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, row):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            r = cell.paragraphs[0].add_run(str(value))
            mark(r)
    return t


def start_addendum(doc, title, subtitle):
    doc.add_page_break()
    add_p(doc, title, style=None, bold=True)
    add_p(doc, subtitle)
    add_p(doc, "黄色高亮为本次 2026-10-11 A/B 批新增或改写内容。Word 打开后请右键目录，选择“更新域 → 更新整个目录”刷新页码。")


def update_05(path):
    doc = Document(path)
    start_addendum(doc, "2026-10-11 Agent 场景与模型白名单补充", "本补充覆盖报表执行、数据分析、库存查询、代码分析和自动五种会话场景，以及后端模型白名单。")
    add_p(doc, "1. 场景不是 Prompt 文案，而是后端固定的工具边界。新建 Agent 会话首条请求传 Scene，后端写入 ZAI_CHAT_SESS.AGENT_TYPE；首条消息后场景固定。空值、GENERAL 和未知值兼容为 AUTO。")
    add_table(doc, ["场景值", "允许工具/行为", "典型问题", "越界行为"], [
        ["REPORT", "报表登记、参数校验、后台作业、READY 结果分析", "执行 ZDPMM702", "表分析请求提示转 DATA"],
        ["DATA", "登记的数据分析和只读 SQL 工具", "按工厂统计 ZIDPMM455", "执行报表提示转 REPORT"],
        ["INVENTORY", "库存候选、快照、收发、批次库存", "A001 在 0001 的库存", "代码/报表请求提示换场景"],
        ["CODE", "read_program_source、表结构读取", "解释 ZDPMM702 程序逻辑", "不执行源码，不读任意对象"],
        ["AUTO", "旧混合路由", "旧 Agent 会话继续提问", "保持历史兼容"],
    ])
    add_p(doc, "2. 工具注册新增 read_program_source。对象名只能是 Z/Y 开头或 /命名空间/；事务码先查 TSTC 主程序；每页最多 400 行、字符上限 18000；模型配置 USER_HIDE_SRC=X 且当前用户不是 ADMIN 时，只返回 TADIR/TRDIRT/修改人等元信息。")
    add_p(doc, "3. Agent 模型白名单的键必须写全：AGENT_MODELS、AGENT_MODELS_RPT、AGENT_MODELS_DATA、AGENT_MODELS_INV、AGENT_MODELS_CODE。场景键优先，空值回退全局，全部为空表示不限制；前端过滤只是体验，后端 agent_model_allowed 才是最终校验。")
    add_p(doc, "4. 普通对话默认不再读取程序源码、表结构或表数据。用户点名对象时，前端显示转 Agent 场景建议；只有 ZAI_CHAT_CFG 的 CHAT_SAP_ACCESS=X 才恢复旧行为。")
    add_p(doc, "5. 能力新增/修改的发布顺序：先贴报表服务和真实作业适配器，再贴对话服务、领域卡、工具注册、意图路由、Agent 服务；SEGW 给 AnalyzeTask 增加 Scene 后 Generate Runtime，最后激活 DPC_EXT。")
    add_p(doc, "维护验收：手动选白名单外模型必须被后端拒绝；报表场景不得调用 DATA/INVENTORY/CODE 工具；CODE 场景必须能看到源码行范围或明确的隐藏源码说明；旧 GENERAL 会话仍按 AUTO。")
    doc.save(path)


def update_06(path):
    doc = Document(path)
    start_addendum(doc, "2026-10-11 A/B 批实施回填", "本补充把原手册中“待修复”的真实作业、结果追问和场景能力更新为当前源码口径。")
    add_p(doc, "1. 当前可部署对象为 9 个整类/整程序替换：ZCL_AI_REPORT_SERVICE V1.0.11、ZCL_AI_REPORT_EXEC_SAP V1.0e、ZCL_AI_CHAT_SERVICE V2.3、ZCL_AI_DOMAIN_CARDS V1.2、ZCL_AI_TOOL_REGISTRY V15.1、ZCL_AI_INTENT_ROUTER V1.0.5、ZCL_AI_AGENT_SERVICE V15.2、ZCL_ZAI_CHAT_DPC_EXT V15.2、ZRMM_IDP_0702 V1.2。线上 ZCL_AI_DATA_ANALYSIS_SERVICE 应为 V13.4c，ZCL_AI_REPORT_MOCK 应为 V1.1，不得用仓库旧版回退。")
    add_table(doc, ["步骤", "SAP 动作", "完成条件"], [
        ["1", "SE24 同时激活 REPORT_SERVICE 与 REPORT_EXEC_SAP", "两类互相引用，均激活；JOB_CLOSE 使用 strtimmed='X'"],
        ["2", "SE24 激活 CHAT_SERVICE", "v8_get_role/read_program_meta 可被 Agent 复用"],
        ["3", "SE24 激活 DOMAIN_CARDS → TOOL_REGISTRY → INTENT_ROUTER", "CODE 域、read_program_source、场景路由可检查"],
        ["4", "SE24 激活 AGENT_SERVICE", "结果上下文、源码分页、白名单方法存在"],
        ["5", "SEGW AnalyzeTask 增加 Scene", "Edm.String、Nullable，Generate Runtime，metadata 可见"],
        ["6", "SE24 激活 DPC_EXT", "AnalyzeTask 透传 FocusedArtifactId 与 Scene"],
        ["7", "SE38 激活 ZRMM_IDP_0702", "compress=abap_false；AI 通道工厂为空时报错"],
    ])
    add_p(doc, "2. 真实报表执行的关键验证：SM37 中作业必须自动释放；无空闲后台进程时按当前时间释放；JOB_OPEN 后失败要删除空作业；P/Z 状态要能解释；结果 schema 必须以 columns 开头并使用所有行键的并集；READY/FAILED 同步回写任务状态和错误原因。")
    add_p(doc, "3. “分析一下这个结果”路径：前端传 FocusedArtifactId；后端验证用户、会话、READY 状态和可读性；找不到聚焦值时用 find_latest_artifact 按会话和 TCode 回退；结果作为上下文注入，不重新执行报表、不新建 DATA 任务。")
    add_p(doc, "4. 前端与 OData：AnalyzeTask 的 Scene 只在新会话首条传入，旧 metadata 缺少该参数时按 AUTO；Agent 模型白名单键为 AGENT_MODELS、AGENT_MODELS_RPT、AGENT_MODELS_DATA、AGENT_MODELS_INV、AGENT_MODELS_CODE；CHAT_SAP_ACCESS=X 才恢复普通对话读 SAP。")
    add_p(doc, "5. 本批不新建 SE11 表、不添加必填字段；部署前核对 ZAI_RPT_PARAM.SEL_KIND、ZAI_RPT_RUN.SAP_JOB_NAME/SAP_JOB_NUMBER、ZAI_RPT_RESULT.RUN_ID、ZAI_CHAT_SESS.AGENT_TYPE。")
    doc.save(path)


def update_07(path):
    doc = Document(path)
    start_addendum(doc, "2026-10-11 A/B 批测试补充", "所有用例同时核对前端提示、OData 返回、SM37、运行表/任务表和 Artifact 归属。")
    add_table(doc, ["用例", "步骤", "通过标准"], [
        ["A-01 作业自动释放", "执行 ZDPMM702，填写 WERKS_I 后提交，查看 SM37", "作业自动为已释放/运行中；不需要人工释放"],
        ["A-02 无空闲后台进程", "制造/等待 CANT_START_IMMEDIATE", "按当前时间释放，最多延迟一个调度周期，不误报提交失败"],
        ["A-03 结果列完整", "读取含 MIH=0、CONFIG 空值的结果", "schema 有 columns；列为全部行键并集，空值列不丢"],
        ["A-04 结果追问", "READY 后说“分析一下这个结果”，再问哪个工厂 MIH 最大", "同一 ArtifactId；Trace 有引用结果；不新建 DATA 任务"],
        ["B-01 场景边界", "REPORT 场景问表分析，DATA 场景执行报表", "均提示换到对应场景，不能跨域调用工具"],
        ["B-02 代码分析", "CODE 场景解释 ZDPMM702", "Trace 有 read_program_source 和行范围；USER_HIDE_SRC 时仅元信息"],
        ["B-03 模型白名单", "只开放一个模型后手动选其他模型", "前端下拉过滤；后端返回“未开放用于 Agent”"],
        ["B-04 普通对话", "对话页询问 ZDPMM702 是什么", "默认不读 SAP；显示转 CODE Agent 卡片；CHAT_SAP_ACCESS=X 才恢复旧行为"],
        ["B-05 旧会话", "打开 AGENT_TYPE 为空或 GENERAL 的旧会话", "按 AUTO 混合路由，历史消息可继续使用"],
    ])
    add_p(doc, "证据要求：每条用例记录 SessionId、TaskId、ReportRunId、ArtifactId、Scene、ModelId、OData 返回、SM37 作业号、ZAI_RPT_RUN 状态和截图/日志时间。只写“通过”而没有这些证据不算放行。")
    add_p(doc, "失败排查优先看：JOB_CLOSE strtimmed 是否为 X；空作业是否被 BP_JOB_DELETE 删除；结果 schema 是否含 columns；UNKNOWN 是否未写 finished_at；AnalyzeTask metadata 是否已生成 Scene；后端是否命中 agent_model_allowed。")
    doc.save(path)


def update_backend_detail(path):
    doc = Document(path)
    start_addendum(doc, "2026-10-11 当前部署入口补充", "原 V13 文档保留历史建表字段；本页是当前 A/B 批实际执行顺序，防止照旧版本误贴源码。")
    add_p(doc, "当前批次不新增 SE11 表、不新增必填字段。整类/整程序替换对象共 9 个：REPORT_SERVICE V1.0.11、REPORT_EXEC_SAP V1.0e、CHAT_SERVICE V2.3、DOMAIN_CARDS V1.2、TOOL_REGISTRY V15.1、INTENT_ROUTER V1.0.5、AGENT_SERVICE V15.2、DPC_EXT V15.2、ZRMM_IDP_0702 V1.2。")
    add_table(doc, ["顺序", "事务", "对象/动作", "必须核对"], [
        ["1", "SE24", "REPORT_SERVICE + REPORT_EXEC_SAP", "一起批量激活；adapter_real_for_run；strtimmed='X'"],
        ["2", "SE24", "CHAT_SERVICE", "v8_get_role/read_program_meta PUBLIC"],
        ["3", "SE24", "DOMAIN_CARDS → TOOL_REGISTRY → INTENT_ROUTER", "CODE 卡、read_program_source、场景约束"],
        ["4", "SE24", "AGENT_SERVICE", "结果上下文、400 行源码页、USER_HIDE_SRC、模型白名单"],
        ["5", "SEGW", "AnalyzeTask 增 Scene", "Edm.String、Nullable；Generate Runtime"],
        ["6", "SE24", "DPC_EXT", "FocusedArtifactId 与 Scene 透传后激活"],
        ["7", "SE38", "ZRMM_IDP_0702", "compress=false；AI 通道空工厂 MESSAGE TYPE E"],
    ])
    add_p(doc, "SEGW AnalyzeTask 的 Scene 可选值只有 REPORT、DATA、INVENTORY、CODE、AUTO。FocusedArtifactId 已存在，不要重复创建。旧前端或旧 metadata 不传 Scene 时，后端按 AUTO 兼容。")
    add_p(doc, "ZAI_CHAT_CFG 可选键：AGENT_MODELS、AGENT_MODELS_RPT、AGENT_MODELS_DATA、AGENT_MODELS_INV、AGENT_MODELS_CODE、CHAT_SAP_ACCESS。场景键覆盖全局键，空值表示不限制；前端设置页保存，后端强制校验。")
    add_p(doc, "报表真实作业验收必须同时检查：SM37 已释放/运行中、无空闲进程降级、失败空作业删除、P/Z 状态解释、columns 列并集、READY/FAILED 任务回写、ArtifactId 归属和结果追问上下文。")
    add_p(doc, "仓库版本落后提醒：ZCL_AI_DATA_ANALYSIS_SERVICE 线上 V13.4c、ZCL_AI_REPORT_MOCK 线上 V1.1；当前部署不要用仓库旧版本覆盖线上。")
    doc.save(path)


def update_historical_highlight(path):
    doc = Document(path)
    start_addendum(doc, "2026-10-11 当前批次索引", "本历史高亮版不再作为当前部署入口；当前部署请转到 docs/03-后端实施文档.docx 和 backend-abap/00_对象清单与部署顺序_20261011.md。")
    add_p(doc, "A/B 批共 9 个对象整类/整程序替换，SEGW 只给 AnalyzeTask 增加可选 Scene（Edm.String、Nullable）并 Generate Runtime；本批不新建表、不加必填字段。")
    add_p(doc, "核心核对：JOB_CLOSE strtimmed='X'、结果 schema columns、compress=abap_false、FocusedArtifactId 结果归属、REPORT/DATA/INVENTORY/CODE/AUTO 场景、read_program_source 400 行分页、Agent 模型白名单。")
    doc.save(path)


TARGETS = {
    "05-Agent能力维护说明.docx": update_05,
    "06-Agent多业务任务与报表分析开发实施与使用说明_V1.0_20261007.docx": update_06,
    "07-Agent多业务任务与报表分析测试执行手册_V1.0_20261008.docx": update_07,
    "SAP_AI对话助手_后端部署详细文档_V13.docx": update_backend_detail,
    "03-后端实施文档_修订高亮版_V13-3.docx": update_historical_highlight,
}


for name, fn in TARGETS.items():
    p = ROOT / "docs" / name
    if any("2026-10-11" in para.text for para in Document(p).paragraphs):
        print(f"skip already updated {p}")
        continue
    fn(p)
    print(f"updated {p}")
