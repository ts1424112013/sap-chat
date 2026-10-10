from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX


ROOT = Path(__file__).resolve().parents[1]


def mark(run, bold=False):
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    run.bold = bold
    return run


def para(doc, text, bold=False):
    p = doc.add_paragraph()
    mark(p.add_run(text), bold)
    return p


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    try:
        t.style = "Table Grid"
    except KeyError:
        pass
    for cell, value in zip(t.rows[0].cells, headers):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        mark(cell.paragraphs[0].add_run(value), True)
    for row in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, row):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            mark(cell.paragraphs[0].add_run(str(value)))
    return t


def addendum(doc, title, intro):
    doc.add_page_break()
    para(doc, title, True)
    para(doc, intro)
    para(doc, "本页是 2026-10-11 Agent 测评增量；黄色高亮为新增内容。Word 打开后请 Ctrl+A → F9 更新目录域。")


def update(path, kind):
    doc = Document(path)
    if any("20261011 Agent 测评" in p.text for p in doc.paragraphs):
        print(f"skip already updated {path}")
        return
    if kind == "maintenance":
        addendum(doc, "20261011 Agent 测评维护说明", "管理员测评随前端 BSP 发布，题库、评分和报告格式由源码固定；它不新增 ABAP 对象，也不绕过后端权限。")
        para(doc, "入口：管理员登录后进入“设置 → 测评 → 打开 Agent 测评”。普通用户没有页签。内置题库 webapp/model/evalCases.json 当前版本 20261011，共 29 题：DATA 12（D01-D12）、REPORT 5（R01-R05）、INVENTORY 3（I01-I03）、CODE 5（C01-C05）、AUTO 4（G01-G04）。")
        para(doc, "AgentEval.js 通过 ChatService.sendTask 使用用户同一条 AnalyzeTask 链路；题目顺序执行，重复次数限制 1~5，多轮共享同一临时会话。requiresScene=true 的题在 metadata 没有 AnalyzeTask.Scene 时跳过；AUTO 不发送 Scene。默认 cleanup=true，题目结束删除会话，跳过题不创建会话。")
        table(doc, ["维护项", "源码口径", "维护要求"], [
            ["题目", "id/name/scene/turns/expect；支持 question 兼容写法，turns 最多 5 轮", "新增题保留场景、版本和变更原因；重复 ID 会自动改名"],
            ["评分", "tools/toolsAny/noTools/toolsNotOk、回答关键词/正则、轮数/耗时、module/needsInput/deterministic", "expect 默认只接受 OK；所有轮次不能 ERROR，最后一轮检查详细预期"],
            ["报告", "Markdown 汇总 + JSON 逐轮明细 + JSON 题库", "保留题库版本、模型、次数、Scene 支持状态和前后端基线"],
            ["安全", "REPORT 只返回参数卡，工具调用仍受 ABAP 白名单/权限约束", "不能用题库改动绕过 SQL、源码或模型白名单"],
        ])
        para(doc, "题库导入格式：{cases:[{scene:'DATA',turns:['问题'],expect:{toolsAny:['run_sql']}}]}。空题或无效 JSON 被拒绝，导入只影响当前浏览器会话，不写 ZAI_CHAT_CFG。演示模式结果是内置示例，只用于检查界面。")
    elif kind == "implementation":
        addendum(doc, "20261011 Agent 测评部署回填", "后端 9 个对象激活后，必须用管理员测评保留部署前后行为基线；测评前端随本批 dist 一起上传。")
        para(doc, "构建时确认 dist/model/evalCases.json、dist/utils/AgentEval.js、dist/view/EvalDialog.fragment.xml 均存在。上传并 Ctrl+F5 后，管理员在“设置”看到“测评”页签；普通用户看不到。题库状态应显示 20261011、29 题和五个场景分布。")
        table(doc, ["阶段", "动作", "必须留存"], [
            ["后端部署前", "Scene 尚未出现时跑全部题；需场景题应标记跳过", "部署前 Markdown/JSON、题库版本、模型、时间"],
            ["后端部署后", "确认 $metadata 有 AnalyzeTask.Scene，再用相同模型/次数重跑", "部署后 Markdown/JSON、通过率、场景支持状态"],
            ["差异检查", "看边界题、工具、安全拒绝、耗时和协议错误", "D12/R05/I03/C05 不跨域成功调用；D08/C04 危险操作被拒"],
            ["清理", "保持“测完删除测试会话”勾选，检查会话历史", "不能残留测试 session；清理失败要记录"],
        ])
        para(doc, "测评不是 SAP 新对象：不创建 SE11 表、不创建 Function Import、不替代 SE24/SEGW 验证。出现 ERROR、协议错、跳过数异常或通过率下降时，先保存 JSON Trace 再排查，不要删题或放宽 expect 掩盖回归。")
    elif kind == "test":
        addendum(doc, "20261011 Agent 测评执行用例", "测试执行手册新增可重复的管理员测评套件；结果以导出的 Markdown 和 JSON 为准。")
        table(doc, ["编号", "执行", "通过判定"], [
            ["E-01 题库", "打开测评，检查版本和场景统计", "20261011 / 29 题；五场景均有题；无效题为 0"],
            ["E-02 Scene 门禁", "旧 metadata 运行全部题，再补 Scene 后重跑", "旧版仅 requiresScene 跳过；新版不再跳过；跳过不计通过"],
            ["E-03 场景边界", "D12/R05/I03/C05", "确定性提示换场景，noTools 成立，工具不能成功执行"],
            ["E-04 REPORT", "R01/R02/R03/R04", "REPORT 模块；需要参数卡；未登记报表拒绝；无结果不乱跑工具"],
            ["E-05 DATA", "D01-D11/G04", "按预期命中 run_sql/sap_data_analysis/search_tables；回答有数据关键词"],
            ["E-06 安全", "D08、C04", "写操作或非自开发源码只能 REJECTED，不能出现成功执行"],
            ["E-07 CODE", "C01-C03", "源码题命中 read_program_source；表结构题命中 describe_table；USER_HIDE_SRC 生效"],
            ["E-08 证据", "导出报告并点击失败行", "MD 有汇总/失败回答；JSON 有每轮状态、工具、耗时、原因"],
        ])
        para(doc, "离线纯逻辑回归：node test-server/agent-eval-tests.js（或 cd test-server && npm run test:eval）。该命令不访问网络、不调用模型，覆盖规范化、评分、跳过、同 session 多轮、错误停止、清理、汇总和 Markdown 报告。")
    elif kind == "backend":
        addendum(doc, "20261011 Agent 测评与后端部署边界", "Agent 测评属于前端验收工具；后端只需提供当前 AnalyzeTask/AnalyzeAgent 协议和 Trace，不为测评单独建表或建接口。")
        para(doc, "管理员测评通过已有 OData 链路验证 REPORT/DATA/INVENTORY/CODE/AUTO 的真实路由、工具白名单、Scene 参数、模型白名单、参数卡和错误协议。Scene 未进 metadata 时，requiresScene 题跳过，其余 AUTO 兼容；这不是后端激活失败，但必须在报告中单独写明。")
        para(doc, "部署前后各跑一次 20261011 题库，比较通过率、跳过数、工具被拒、错误/超时/协议错、平均轮数和耗时；报告和 JSON 明细归档。REPORT 题只拿参数卡，不提交真实作业；演示模式不能作为现场通过证据。")
    elif kind == "historical":
        addendum(doc, "20261011 Agent 测评索引", "本历史高亮版仅作交叉参考；当前测评入口和部署步骤以 03/04 文档为准。")
        para(doc, "新增前端文件：model/evalCases.json（20261011、29 题）、utils/AgentEval.js（规范化/顺序执行/评分/汇总/Markdown）、view/EvalDialog.fragment.xml（管理员入口）。不新增 ABAP 类、SE11 表、SEGW 接口。")
        para(doc, "现场证据必须包含 Scene metadata 状态、题库版本、模型/重复次数、跳过数、通过率、工具 Trace、错误/超时/协议错、平均/最长耗时和导出的 MD/JSON；不要只在 Word 中写“测评通过”。")
    doc.save(path)
    print(f"updated {path}")


TARGETS = [
    ("05-Agent能力维护说明.docx", "maintenance"),
    ("06-Agent多业务任务与报表分析开发实施与使用说明_V1.0_20261007.docx", "implementation"),
    ("07-Agent多业务任务与报表分析测试执行手册_V1.0_20261008.docx", "test"),
    ("SAP_AI对话助手_后端部署详细文档_V13.docx", "backend"),
    ("03-后端实施文档_修订高亮版_V13-3.docx", "historical"),
]

for filename, kind in TARGETS:
    update(ROOT / "docs" / filename, kind)
